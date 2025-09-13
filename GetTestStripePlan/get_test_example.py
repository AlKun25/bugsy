#!/usr/bin/env python3
from __future__ import annotations

import os
import sys
import json
import argparse
from typing import List, Optional

from pydantic import BaseModel, Field, ValidationError

# ---------- OpenAI SDK ----------
try:
    from openai import OpenAI
except Exception:
    print("Please install the OpenAI SDK: pip install --upgrade openai", file=sys.stderr)
    raise

# ---------- Optional .docx support ----------
try:
    from docx import Document
except Exception:
    Document = None  # still allow .txt flows without python-docx


# ==============================
# Config
# ==============================
MODEL = os.getenv("OPENAI_MODEL", "gpt-4.1-mini")


# ==============================
# Schema for structured outputs
# ==============================
class BugSolution(BaseModel):
    root_cause: str = Field(..., description="One-paragraph explanation of the underlying cause.")
    fix_summary: str = Field(..., description="Short summary (1–2 sentences) of the fix approach.")
    steps_to_reproduce: List[str] = Field(..., description="Exact steps to reproduce the issue.")
    steps_to_fix: List[str] = Field(..., description="Exact steps to implement/verify the fix.")
    code_patch: Optional[str] = Field(None, description="Minimal code patch/snippet if applicable.")
    worked_example_output: Optional[str] = Field(
        None, description="Concrete expected output after fix."
    )
    caveats: Optional[List[str]] = Field(
        default=None, description="Risks, edge cases, or follow-ups."
    )
    confidence: float = Field(..., ge=0.0, le=1.0, description="Model confidence 0..1")


# ==============================
# Utilities
# ==============================
def load_text_file(path: str) -> str:
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def load_prompt_template(prompt_file: str) -> str:
    """
    Load a prompt template from .txt or .docx.
    The template should contain {bug_title} and {bug_description} placeholders.
    """
    if not os.path.exists(prompt_file):
        raise FileNotFoundError(f"Prompt file not found: {prompt_file}")

    if prompt_file.lower().endswith(".txt"):
        return load_text_file(prompt_file)

    if prompt_file.lower().endswith(".docx"):
        if Document is None:
            raise RuntimeError(
                "python-docx is not installed. Run: pip install python-docx"
            )
        doc = Document(prompt_file)
        return "\n".join(p.text for p in doc.paragraphs).strip()

    raise ValueError("Prompt file must be .txt or .docx")


def parse_bug_file(bug_file: str) -> tuple[str, str]:
    """
    Interpret the first non-empty line as title; remaining non-empty lines as description.
    """
    raw = load_text_file(bug_file)
    lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    if not lines:
        raise ValueError("Bug file is empty.")
    title = lines[0]
    description = "\n".join(lines[1:]) if len(lines) > 1 else ""
    return title, description


def build_prompt(bug_title: str, bug_description: str, prompt_file: str) -> str:
    """
    Load a prompt template and fill placeholders literally (no str.format)
    to avoid conflicts with JSON braces in your template.
    """
    template = load_prompt_template(prompt_file)
    return (template
            .replace("{bug_title}", bug_title)
            .replace("{bug_description}", bug_description))


def _extract_text_from_responses(resp) -> str:
    """
    Defensive extraction from Responses API output.
    """
    try:
        return resp.output[0].content[0].text
    except Exception:
        # Last resort: dump the whole output structure
        return json.dumps(getattr(resp, "output", []), default=lambda o: getattr(o, "__dict__", str(o)))


def solve_with_structured_outputs(client: OpenAI, prompt: str) -> BugSolution:
    """
    Preferred: Responses API + structured parse (if available in SDK).
    Fallback: request strict JSON and validate with Pydantic.
    """
    parse = getattr(client.responses, "parse", None)
    if callable(parse):
        return client.responses.parse(
            model=MODEL,
            input=prompt,
            response_format=BugSolution,  # Pydantic class
        )

    # Fallback to JSON-mode
    resp = client.responses.create(
        model=MODEL,
        input=(
            prompt
            + "\n\nReturn ONLY valid JSON with keys: "
              "root_cause, fix_summary, steps_to_reproduce, steps_to_fix, "
              "code_patch, worked_example_output, caveats, confidence."
        ),
        text={"format": {"type": "json_object"}},
    )

    raw = _extract_text_from_responses(resp)

    # Try direct JSON first
    try:
        data = json.loads(raw)
    except Exception:
        # Attempt to carve a JSON object if the model added extraneous text
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end != -1 and end > start:
            data = json.loads(raw[start:end + 1])
        else:
            raise RuntimeError(f"Model did not return JSON. Raw:\n{raw}")

    try:
        return BugSolution(**data)
    except ValidationError as ve:
        raise RuntimeError(
            "Model JSON didn't match schema:\n"
            f"{ve}\n\nRaw JSON:\n{json.dumps(data, indent=2, ensure_ascii=False)}"
        )


# ==============================
# CLI
# ==============================
def main():
    parser = argparse.ArgumentParser(
        description="Generate a bug fix plan with OpenAI using a prompt template and a bug input file."
    )
    parser.add_argument(
        "--bug_file",
        required=True,
        help="Path to a text file containing the bug. First non-empty line = title; remainder = description.",
    )
    parser.add_argument(
        "--prompt_file",
        required=True,
        help="Path to a .txt or .docx prompt template with {bug_title} and {bug_description} placeholders.",
    )
    parser.add_argument(
        "--pretty", action="store_true", help="Pretty-print JSON output"
    )
    args = parser.parse_args()

    # Read inputs
    bug_title, bug_description = parse_bug_file(args.bug_file)
    prompt = build_prompt(bug_title, bug_description, prompt_file=args.prompt_file)

    # OpenAI client
    api_key = os.environ.get("OPENAI_API_KEY")
    if not api_key:
        print("Missing OPENAI_API_KEY environment variable.", file=sys.stderr)
        sys.exit(1)
    client = OpenAI(api_key=api_key)

    # Call model
    solution = solve_with_structured_outputs(client, prompt)

    # Emit JSON
    out = json.loads(solution.model_dump_json())
    print(json.dumps(out, indent=2 if args.pretty else None, ensure_ascii=False))


if __name__ == "__main__":
    main()

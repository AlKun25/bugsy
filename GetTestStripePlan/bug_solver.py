#!/usr/bin/env python3
from __future__ import annotations

import os
import json
from typing import List, Optional

from pydantic import BaseModel, Field, ValidationError

# ---------- OpenAI SDK ----------
try:
    from openai import OpenAI
except Exception:
    OpenAI = None  # Will be checked when actually needed

# ---------- Optional .docx support ----------
try:
    from docx import Document
except Exception:
    Document = None  # still allow .txt flows without python-docx


# ==============================
# Config
# ==============================
MODEL = os.getenv("OPENAI_MODEL", "gpt-4.1-mini")


# ==============================
# Schema for structured outputs
# ==============================
class TestStep(BaseModel):
    """Individual test step."""
    type: str = Field(description="Type of step: 'action' or 'assertion'")
    description: str = Field(description="Description of the step")

class TestCase(BaseModel):
    """Individual test case."""
    id: str = Field(description="Test case ID")
    title: str = Field(description="Test case title")
    description: str = Field(description="Test case description")
    category: str = Field(description="Test category")
    priority: str = Field(description="Test priority")
    steps: List[TestStep] = Field(description="Test steps")

class TestPlan(BaseModel):
    """Complete test plan with multiple test cases."""
    test_cases: List[TestCase] = Field(description="List of test cases")


# ==============================
# File I/O Utilities
# ==============================
def load_text_file(path: str) -> str:
    """Load content from a text file."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def load_prompt_template(prompt_file: str) -> str:
    """
    Load a prompt template from .txt or .docx.
    The template should contain {bug_title} and {bug_description} placeholders.
    """
    if not os.path.exists(prompt_file):
        raise FileNotFoundError(f"Prompt file not found: {prompt_file}")

    if prompt_file.lower().endswith(".txt"):
        return load_text_file(prompt_file)

    if prompt_file.lower().endswith(".docx"):
        if Document is None:
            raise RuntimeError(
                "python-docx is not installed. Run: pip install python-docx"
            )
        doc = Document(prompt_file)
        return "\n".join(p.text for p in doc.paragraphs).strip()

    raise ValueError("Prompt file must be .txt or .docx")


def parse_bug_file(bug_file: str) -> tuple[str, str]:
    """
    Interpret the first non-empty line as title; remaining non-empty lines as description.
    """
    raw = load_text_file(bug_file)
    lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    if not lines:
        raise ValueError("Bug file is empty.")
    title = lines[0]
    description = "\n".join(lines[1:]) if len(lines) > 1 else ""
    return title, description


# ==============================
# Prompt Building
# ==============================
def build_prompt(bug_title: str, bug_description: str, prompt_file: str) -> str:
    """
    Load a prompt template and fill placeholders literally (no str.format)
    to avoid conflicts with JSON braces in your template.
    """
    template = load_prompt_template(prompt_file)
    return (template
            .replace("{bug_title}", bug_title)
            .replace("{bug_description}", bug_description))


# ==============================
# OpenAI Integration
# ==============================
def solve_with_structured_outputs(client: OpenAI, prompt: str) -> TestPlan:
    """
    Generate test cases using OpenAI's structured outputs.
    """
    # Enhanced prompt for test case generation
    test_prompt = f"""{prompt}

Generate exactly 16 comprehensive test cases in the following format. Each test case should have:
- A unique ID (TC001, TC002, etc.)
- A descriptive title
- A detailed description
- A category (functional, error handling, security, performance, etc.)
- A priority (High, Medium, Low)
- Multiple detailed steps with type 'action' or 'assertion'

Ensure the test cases cover:
1. Happy path scenarios
2. Error handling and edge cases
3. Security considerations
4. User authentication flows
5. Data validation
6. UI/UX interactions
7. Performance considerations
8. Integration scenarios

Make sure each test case has at least 4-6 detailed steps."""
    
    try:
        # Use the current OpenAI API with structured outputs
        response = client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "user", "content": test_prompt}
            ],
            response_format={
                "type": "json_schema",
                "json_schema": {
                    "name": "test_plan",
                    "schema": TestPlan.model_json_schema()
                }
            }
        )
        
        raw_content = response.choices[0].message.content
        data = json.loads(raw_content)
        return TestPlan(**data)
        
    except Exception as e:
        # Fallback to regular JSON mode
        response = client.chat.completions.create(
            model=MODEL,
            messages=[
                {
                    "role": "user", 
                    "content": test_prompt + "\n\nReturn ONLY valid JSON with a 'test_cases' array containing test case objects."
                }
            ],
            response_format={"type": "json_object"}
        )
        
        raw_content = response.choices[0].message.content
        
        # Try direct JSON first
        try:
            data = json.loads(raw_content)
        except Exception:
            # Attempt to carve a JSON object if the model added extraneous text
            start, end = raw_content.find("{"), raw_content.rfind("}")
            if start != -1 and end != -1 and end > start:
                data = json.loads(raw_content[start:end + 1])
            else:
                raise RuntimeError(f"Model did not return JSON. Raw:\n{raw_content}")
        
        try:
            return TestPlan(**data)
        except ValidationError as ve:
            raise RuntimeError(
                "Model JSON didn't match schema:\n"
                f"{ve}\n\nRaw JSON:\n{json.dumps(data, indent=2, ensure_ascii=False)}"
            )


# ==============================
# Main Processing Function
# ==============================
def process_bug_report(bug_content: str, prompt_content: str) -> TestPlan:
    """
    Main processing function that combines bug report and prompt,
    then calls OpenAI for test case generation.
    """
    if OpenAI is None:
        raise ImportError(
            "OpenAI SDK not available. Install with: pip install openai"
        )
    
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise ValueError("Missing OPENAI_API_KEY environment variable")
    
    client = OpenAI(api_key=api_key)
    
    # Combine the inputs into a comprehensive prompt
    full_prompt = f"""
{prompt_content}

Bug Report:
{bug_content}

Please generate comprehensive test cases based on this bug report and requirements.
"""
    
    return solve_with_structured_outputs(client, full_prompt)